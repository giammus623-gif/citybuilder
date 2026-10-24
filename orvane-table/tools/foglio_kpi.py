"""Crea il foglio di calcolo del fondatore: contatti con punteggio (Parte 3.6 e 8.6),
cruscotto settimanale e mensile (Parte 15), clienti, concorrenti e listino.

    npm run foglio            ->  docs/fogli/orvane-cruscotto.xlsx

Tutti i calcoli sono formule: il foglio si aggiorna da solo quando si scrivono i numeri.
"""

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DESTINAZIONE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "docs" / "fogli" / "orvane-cruscotto.xlsx"

FONT = "Arial"
BLU = "0000FF"
VERDE = "008000"
F_TESTO = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color=BLU)
F_LINK = Font(name=FONT, size=10, color=VERDE)
F_TITOLO = Font(name=FONT, size=14, bold=True)
F_INTESTAZIONE = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_NOTA = Font(name=FONT, size=9, italic=True, color="555555")
SFONDO_INTESTAZIONE = PatternFill("solid", fgColor="8A2D1C")
DA_COMPILARE = PatternFill("solid", fgColor="FFF2A8")
ROSSO = PatternFill("solid", fgColor="F8D0CC")
AMBRA = PatternFill("solid", fgColor="FCE8C3")
VERDINO = PatternFill("solid", fgColor="D9EFD8")
GRIGIO = PatternFill("solid", fgColor="E8E8E8")
LINEA = Border(bottom=Side(style="thin", color="DDDDDD"))
EURO = '#,##0 "€";-#,##0 "€";"-"'
EURO2 = '#,##0.00 "€";-#,##0.00 "€";"-"'
PERC = '0.0%;-0.0%;"-"'
DATA = "DD/MM/YYYY"

wb = Workbook()


def intestazione(ws, riga, titoli, larghezze=None):
    for i, t in enumerate(titoli, start=1):
        c = ws.cell(row=riga, column=i, value=t)
        c.font = F_INTESTAZIONE
        c.fill = SFONDO_INTESTAZIONE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[riga].height = 32
    if larghezze:
        for i, w in enumerate(larghezze, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=riga + 1, column=1)


def titolo(ws, testo, nota):
    ws["A1"] = testo
    ws["A1"].font = F_TITOLO
    ws["A2"] = nota
    ws["A2"].font = F_NOTA


def scrivi(ws, rif, valore, font=F_TESTO, formato=None, riempimento=None):
    c = ws[rif]
    c.value = valore
    c.font = font
    if formato:
        c.number_format = formato
    if riempimento:
        c.fill = riempimento
    return c


# ---------------------------------------------------------------------------
# Leggimi

ws = wb.active
ws.title = "Leggimi"
titolo(ws, "Orvane Table · cruscotto del fondatore", "Generato da tools/foglio_kpi.py (npm run foglio). Riferimenti: Parti 3.6, 8.6 e 15 del prompt master.")
righe = [
    ("Come si usa", ""),
    ("Celle gialle", "da compilare: numeri della settimana, del mese, contatti e clienti"),
    ("Testo blu", "valori inseriti a mano: obiettivi, soglie, listino. Si possono cambiare"),
    ("Testo nero", "formule: non scriverci sopra"),
    ("Testo verde", "valori presi da un altro foglio"),
    ("Celle rosse", "sotto soglia: si corregge prima l'indicatore più a sinistra (Parte 15.5)"),
    ("", ""),
    ("Ogni venerdì (10 minuti)", "riga della settimana in Settimana; prima cella rossa da sinistra; prompt 17.14 se serve"),
    ("Ogni sera", "Contatti: filtrare la colonna Scaduta e fare le azioni del giorno"),
    ("Primo venerdì del mese", "riga del mese in Mese; se la differenza dal modello supera il 30%, aggiornare il modello (15.5)"),
    ("", ""),
    ("Regole di lettura", "una settimana bassa è rumore, due di fila sono un segnale; nessun cambio di prezzo con meno di 10 osservazioni"),
    ("Righe di esempio", "la prima riga di Contatti e di Clienti è un esempio con dati inventati: cancellatela quando inserite i vostri"),
]
for i, (a, b) in enumerate(righe, start=4):
    scrivi(ws, f"A{i}", a, Font(name=FONT, size=10, bold=bool(a and not b)))
    scrivi(ws, f"B{i}", b)
for r, fill in ((5, DA_COMPILARE), (9, ROSSO)):
    ws[f"A{r}"].fill = fill
ws["A6"].font = F_INPUT
ws["A8"].font = F_LINK
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110

# ---------------------------------------------------------------------------
# Obiettivi e soglie (Parte 15.1 e 10.5)

ob = wb.create_sheet("Obiettivi")
titolo(ob, "Obiettivi e soglie", "Valori della Parte 15.1 e del modello della Parte 10 (scenario realistico). Si possono cambiare: tutte le formule li leggono da qui.")
intestazione(ob, 4, ["Indicatore settimanale", "Obiettivo", "Soglia di allarme", "Azione se sotto soglia"], [34, 14, 16, 60])
settimanali = [
    ("Contatti qualificati", 14, 8, "Bloccare fasce orarie in agenda, allargare la lista"),
    ("Conversazioni", 4, 2, "Cambiare apertura (Parte 9.1)"),
    ("Demo", 2, 1, "Rendere la demo più veloce: pronta in 24 ore"),
    ("Proposte", 1, 0.5, "Rivedere la qualificazione (allarme: 0 per due settimane)"),
    ("Ore di vendita (massimo)", 16, 20, "Oltre la soglia per tre settimane: più segnalazioni e partner"),
    ("Conversione contatti-contratti", 0.045, 0.02, "Sotto il 2% su 60 contatti: analisi dei motivi di perdita"),
]
for i, (nome, obiettivo, soglia, azione) in enumerate(settimanali, start=5):
    scrivi(ob, f"A{i}", nome)
    formato = PERC if nome.startswith("Conversione") else "0.0"
    scrivi(ob, f"B{i}", obiettivo, F_INPUT, formato)
    scrivi(ob, f"C{i}", soglia, F_INPUT, formato)
    scrivi(ob, f"D{i}", azione)
ob["C8"].comment = Comment("La Parte 15.1 dà come allarme 0 proposte per due settimane: qui la soglia 0,5 colora di rosso lo zero.", "Orvane")

intestazione_mese = 13
for i, t in enumerate(["Mese", "Periodo", "Contratti firmati", "Clienti attivi a fine mese", "Ricorrente incassato (€)", "Cassa cumulata (€)"], start=1):
    c = ob.cell(row=intestazione_mese, column=i, value=t)
    c.font = F_INTESTAZIONE
    c.fill = SFONDO_INTESTAZIONE
    c.alignment = Alignment(wrap_text=True)
ob.row_dimensions[intestazione_mese].height = 32
scrivi(ob, "A12", "Modello di riferimento, scenario realistico (Parte 10.5)", Font(name=FONT, size=10, bold=True))
modello = [
    (0, 0, 0, 0),
    (2, 2, 0, 158),
    (1, 3, 221, 544),
    (2, 4.9, 322, 1221),
    (2, 6.8, 526, 2361),
    (2, 8.6, 725, 3596),
    (3, 11.3, 917, 5126),
    (3, 14.0, 1212, 7041),
    (3, 16.5, 1497, 9092),
    (3, 19.0, 1773, 11275),
    (3, 21.5, 2042, 13587),
    (3, 23.8, 2302, 15964),
]
mesi_nomi = ["nov 2026", "dic 2026", "gen 2027", "feb 2027", "mar 2027", "apr 2027", "mag 2027", "giu 2027", "lug 2027", "ago 2027", "set 2027", "ott 2027"]
for i, (contratti, clienti, ricorrente, cassa) in enumerate(modello):
    r = intestazione_mese + 1 + i
    scrivi(ob, f"A{r}", i + 1)
    scrivi(ob, f"B{r}", mesi_nomi[i])
    scrivi(ob, f"C{r}", contratti, F_INPUT)
    scrivi(ob, f"D{r}", clienti, F_INPUT, "0.0")
    scrivi(ob, f"E{r}", ricorrente, F_INPUT, EURO)
    scrivi(ob, f"F{r}", cassa, F_INPUT, EURO)
ob[f"A{intestazione_mese + 13}"] = "Fonte: tabella 10.5 del prompt master; ricalcolabile con npm run modello. Se il primo incasso slitta, spostare i periodi di un mese (Parte 8.9)."
ob[f"A{intestazione_mese + 13}"].font = F_NOTA
ob.column_dimensions["E"].width = 16
ob.column_dimensions["F"].width = 16

# ---------------------------------------------------------------------------
# Listino (Parte 4.1 e 4.2)

li = wb.create_sheet("Listino")
titolo(li, "Listino", "Prezzi indicativi al netto di IVA (Parte 4.1), da validare nelle prime dieci conversazioni. Il canone non si abbassa mai (4.6).")
intestazione(li, 4, ["Pacchetto", "Setup (€)", "Canone mensile (€)", "Costo variabile stimato (€/mese)", "Minuti inclusi", "Margine mensile (€)", "Margine %"], [16, 12, 16, 20, 14, 16, 12])
for i, (nome, setup, canone, variabile, minuti) in enumerate(
    [("Start", 290, 49, 3, 0), ("Pro", 790, 129, 30, 250), ("Premium", 1490, 219, 65, 500)], start=5
):
    scrivi(li, f"A{i}", nome)
    scrivi(li, f"B{i}", setup, F_INPUT, EURO)
    scrivi(li, f"C{i}", canone, F_INPUT, EURO)
    scrivi(li, f"D{i}", variabile, F_INPUT, EURO)
    scrivi(li, f"E{i}", minuti, F_INPUT, "0")
    scrivi(li, f"F{i}", f"=C{i}-D{i}", formato=EURO)
    scrivi(li, f"G{i}", f"=IFERROR(F{i}/C{i},0)", formato=PERC)
scrivi(li, "A9", "Minuto extra (€)")
scrivi(li, "B9", 0.25, F_INPUT, EURO2)
scrivi(li, "A10", "Sconto cliente fondatore")
scrivi(li, "B10", 0.5, F_INPUT, PERC)
scrivi(li, "A11", "Fonte: Parte 4 del prompt master; stesso listino di src/modello/listino.ts.", F_NOTA)

# ---------------------------------------------------------------------------
# Contatti (Parte 8.6 + scheda di qualificazione 3.6)

ct = wb.create_sheet("Contatti")
titolo(ct, "Contatti", "Una riga per locale. Criteri: 1 = vero, 0 = falso, vuoto = da verificare. Punteggio ≥ 7: demo entro una settimana; 4-6: un messaggio al mese; sotto 4: archiviare (Parte 3.6).")
criteri = [
    "Titolare decide da solo",
    "Almeno 30 coperti",
    "Prenotazioni per telefono",
    "Ammette chiamate perse",
    "Sito o scheda Google trascurati",
    "Usa WhatsApp per prenotare",
    "Più di 20 coperti a sera",
    "Tavoli o sale per gruppi",
    "Aperto da più di 3 anni",
    "Contatto personale",
]
colonne = (
    ["Locale", "Indirizzo", "Telefono", "Titolare", "Segmento"]
    + criteri
    + ["Punteggio", "Da verificare", "Cosa fare", "Primo contatto", "Canale", "Esito", "Prossima azione", "Data prossima azione", "Scaduta", "Pacchetto proposto", "Canone proposto (€)", "Note", "Motivo della perdita"]
)
larghezze = [24, 26, 15, 14, 18] + [9] * 10 + [10, 10, 26, 13, 12, 16, 28, 13, 9, 12, 12, 30, 22]
intestazione(ct, 4, colonne, larghezze)
ct.row_dimensions[4].height = 60
C_PUNTI, C_VERIFICA, C_AZIONE, C_DATA_PROSSIMA, C_SCADUTA, C_PACCHETTO, C_CANONE = 16, 17, 18, 23, 24, 25, 26
PRIMA, ULTIMA = 5, 104
lettera = get_column_letter
for r in range(PRIMA, ULTIMA + 1):
    for col in list(range(1, 16)) + [19, 20, 21, 22, C_DATA_PROSSIMA, C_PACCHETTO, 27, 28]:
        ct.cell(row=r, column=col).fill = DA_COMPILARE
        ct.cell(row=r, column=col).font = F_TESTO
    rng = f"F{r}:O{r}"
    ct.cell(row=r, column=C_PUNTI, value=f'=IF(COUNTA(A{r})=0,"",COUNTIF({rng},1))').font = F_TESTO
    ct.cell(row=r, column=C_VERIFICA, value=f'=IF(COUNTA(A{r})=0,"",COUNTBLANK({rng}))').font = F_TESTO
    p = f"{lettera(C_PUNTI)}{r}"
    ct.cell(row=r, column=C_AZIONE, value=f'=IF({p}="","",IF({p}>=7,"Demo entro una settimana",IF({p}>=4,"Un messaggio al mese","Archiviare")))').font = F_TESTO
    d = f"{lettera(C_DATA_PROSSIMA)}{r}"
    ct.cell(row=r, column=C_SCADUTA, value=f'=IF(AND(ISNUMBER({d}),{d}<TODAY()),"SÌ","")').font = F_TESTO
    pk = f"{lettera(C_PACCHETTO)}{r}"
    ct.cell(row=r, column=C_CANONE, value=f'=IF({pk}="","",IFERROR(INDEX(Listino!$C$5:$C$7,MATCH({pk},Listino!$A$5:$A$7,0)),""))').font = F_LINK
    ct.cell(row=r, column=C_CANONE).number_format = EURO
    for col in (19, C_DATA_PROSSIMA):
        ct.cell(row=r, column=col).number_format = DATA

esempio = ["Trattoria Esempio (riga di prova)", "Via Roma 1, Città", "+39 0000 000000", "Mario", "Ristorante 30-120"] + [1, 1, 1, 1, None, 1, 1, 0, 1, 0]
for col, v in enumerate(esempio, start=1):
    ct.cell(row=PRIMA, column=col, value=v)
ct.cell(row=PRIMA, column=19, value=dt.date(2026, 11, 4))
ct.cell(row=PRIMA, column=20, value="Visita")
ct.cell(row=PRIMA, column=21, value="Conversazione")
ct.cell(row=PRIMA, column=22, value="Preparare demo sul suo menu")
ct.cell(row=PRIMA, column=C_DATA_PROSSIMA, value=dt.date(2026, 11, 6))
ct.cell(row=PRIMA, column=C_PACCHETTO, value="Pro")
ct.cell(row=PRIMA, column=27, value='Perde "4-5 chiamate il sabato sera" (parole del titolare)')

dv_segmento = DataValidation(type="list", formula1='"Ristorante 30-120,Pizzeria,Agriturismo,Sala eventi,Bar con cucina,Da evitare"', allow_blank=True)
dv_criterio = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=True, error="Scrivere 1 (vero), 0 (falso) o lasciare vuoto (da verificare)")
dv_canale = DataValidation(type="list", formula1='"Visita,Telefonata,WhatsApp,Rete personale,Segnalazione,Partner,Video"', allow_blank=True)
dv_esito = DataValidation(type="list", formula1='"Da contattare,Contattato,Conversazione,Demo,Proposta,Firmato,Perso,Archiviato"', allow_blank=True)
dv_pacchetto = DataValidation(type="list", formula1='"Start,Pro,Premium"', allow_blank=True)
for dv, area in (
    (dv_segmento, f"E{PRIMA}:E{ULTIMA}"),
    (dv_criterio, f"F{PRIMA}:O{ULTIMA}"),
    (dv_canale, f"T{PRIMA}:T{ULTIMA}"),
    (dv_esito, f"U{PRIMA}:U{ULTIMA}"),
    (dv_pacchetto, f"{lettera(C_PACCHETTO)}{PRIMA}:{lettera(C_PACCHETTO)}{ULTIMA}"),
):
    ct.add_data_validation(dv)
    dv.add(area)
punti = f"{lettera(C_PUNTI)}{PRIMA}:{lettera(C_PUNTI)}{ULTIMA}"
ct.conditional_formatting.add(punti, FormulaRule(formula=[f'AND(ISNUMBER({lettera(C_PUNTI)}{PRIMA}),{lettera(C_PUNTI)}{PRIMA}>=7)'], fill=VERDINO))
ct.conditional_formatting.add(punti, FormulaRule(formula=[f'AND(ISNUMBER({lettera(C_PUNTI)}{PRIMA}),{lettera(C_PUNTI)}{PRIMA}>=4,{lettera(C_PUNTI)}{PRIMA}<7)'], fill=AMBRA))
ct.conditional_formatting.add(punti, FormulaRule(formula=[f'AND(ISNUMBER({lettera(C_PUNTI)}{PRIMA}),{lettera(C_PUNTI)}{PRIMA}<4)'], fill=GRIGIO))
ct.conditional_formatting.add(f"{lettera(C_SCADUTA)}{PRIMA}:{lettera(C_SCADUTA)}{ULTIMA}", CellIsRule(operator="equal", formula=['"SÌ"'], fill=ROSSO))
ct.auto_filter.ref = f"A4:{lettera(len(colonne))}{ULTIMA}"

# ---------------------------------------------------------------------------
# Settimana (Parte 15.1)

st = wb.create_sheet("Settimana")
titolo(st, "Settimana", "Una riga per settimana, compilata il venerdì. Rosso = sotto soglia; SEGNALE = due settimane consecutive sotto soglia nei contatti (Parte 15.5).")
intestazione(
    st,
    4,
    ["N.", "Dal (lunedì)", "Contatti", "Conversazioni", "Demo", "Proposte", "Contratti", "Ore di vendita", "Cassa (€)", "Conversione cumulata", "Contatti cumulati", "Segnale contatti", "Note: cosa è andato storto, tre obiettivi"],
    [5, 12, 10, 13, 8, 10, 10, 12, 12, 13, 12, 12, 60],
)
INIZIO = dt.date(2026, 10, 26)
R0, RN = 5, 56
for i, r in enumerate(range(R0, RN + 1)):
    scrivi(st, f"A{r}", i + 1)
    scrivi(st, f"B{r}", INIZIO + dt.timedelta(weeks=i), formato=DATA)
    for col in "CDEFGHI":
        c = st[f"{col}{r}"]
        c.fill = DA_COMPILARE
        c.font = F_TESTO
    st[f"I{r}"].number_format = EURO
    st[f"M{r}"].fill = DA_COMPILARE
    st[f"J{r}"] = f'=IF(SUM($C${R0}:C{r})=0,"",SUM($G${R0}:G{r})/SUM($C${R0}:C{r}))'
    st[f"J{r}"].number_format = PERC
    st[f"K{r}"] = f'=IF(COUNT(C{r})=0,"",SUM($C${R0}:C{r}))'
    precedente = f"C{r - 1}" if r > R0 else None
    st[f"L{r}"] = (
        f'=IF(AND(COUNT(C{r})=1,COUNT({precedente})=1),IF(AND(C{r}<Obiettivi!$C$5,{precedente}<Obiettivi!$C$5),"SEGNALE",""),"")'
        if precedente
        else '=""'
    )
    for col in "JKL":
        st[f"{col}{r}"].font = F_TESTO
for col, riga_obiettivo in (("C", 5), ("D", 6), ("E", 7), ("F", 8)):
    st.conditional_formatting.add(
        f"{col}{R0}:{col}{RN}",
        FormulaRule(formula=[f"AND(ISNUMBER({col}{R0}),{col}{R0}<Obiettivi!$C${riga_obiettivo})"], fill=ROSSO),
    )
st.conditional_formatting.add(f"H{R0}:H{RN}", FormulaRule(formula=[f"AND(ISNUMBER(H{R0}),H{R0}>Obiettivi!$C$9)"], fill=ROSSO))
st.conditional_formatting.add(f"J{R0}:J{RN}", FormulaRule(formula=[f"AND(ISNUMBER(J{R0}),K{R0}>=60,J{R0}<Obiettivi!$C$10)"], fill=ROSSO))
st.conditional_formatting.add(f"L{R0}:L{RN}", CellIsRule(operator="equal", formula=['"SEGNALE"'], fill=ROSSO))

# ---------------------------------------------------------------------------
# Mese (Parte 15.2 e 15.3)

ms = wb.create_sheet("Mese")
titolo(ms, "Mese", "Primo venerdì del mese. Obiettivi dal foglio Obiettivi; rosso = sotto il 70% dell'obiettivo o oltre la soglia della Parte 15.2-15.3.")
titoli_mese = [
    "Mese", "Periodo", "Clienti a inizio mese", "Clienti persi", "Clienti a fine mese", "Ricorrente incassato (€)", "Costi variabili (€)",
    "Cassa cumulata (€)", "Fatture scadute > 15 gg", "Ricavo medio per cliente", "Margine lordo", "Abbandono mensile",
    "Clienti: % obiettivo", "Ricorrente: % obiettivo", "Cassa: % obiettivo", "Risolte senza persona", "Domande aperte",
    "Ore medie di correzione", "Minuti usati / inclusi", "Giorni di attivazione (media)", "Soddisfazione (0-10)",
]
intestazione(ms, 4, titoli_mese, [6, 10, 11, 9, 11, 13, 12, 13, 11, 12, 10, 11, 11, 11, 11, 11, 10, 11, 11, 12, 11])
ms.row_dimensions[4].height = 60
for i in range(12):
    r = 5 + i
    ro = 14 + i  # riga del modello nel foglio Obiettivi
    scrivi(ms, f"A{r}", i + 1)
    scrivi(ms, f"B{r}", mesi_nomi[i])
    for col in "CDEFGHIPQRSTU":
        ms[f"{col}{r}"].fill = DA_COMPILARE
        ms[f"{col}{r}"].font = F_TESTO
    for col in "FGH":
        ms[f"{col}{r}"].number_format = EURO
    for col in "PS":
        ms[f"{col}{r}"].number_format = PERC
    ms[f"J{r}"] = f'=IF(OR(COUNT(F{r})=0,COUNT(E{r})=0,E{r}=0),"",F{r}/E{r})'
    ms[f"J{r}"].number_format = EURO2
    ms[f"K{r}"] = f'=IF(OR(COUNT(F{r})=0,F{r}=0),"",(F{r}-G{r})/F{r})'
    ms[f"L{r}"] = f'=IF(OR(COUNT(C{r})=0,C{r}=0,COUNT(D{r})=0),"",D{r}/C{r})'
    ms[f"M{r}"] = f'=IF(OR(COUNT(E{r})=0,Obiettivi!$D${ro}=0),"",E{r}/Obiettivi!$D${ro})'
    ms[f"N{r}"] = f'=IF(OR(COUNT(F{r})=0,Obiettivi!$E${ro}=0),"",F{r}/Obiettivi!$E${ro})'
    ms[f"O{r}"] = f'=IF(OR(COUNT(H{r})=0,Obiettivi!$F${ro}=0),"",H{r}/Obiettivi!$F${ro})'
    for col in "KLMNO":
        ms[f"{col}{r}"].number_format = PERC
    for col in "JKLMNO":
        ms[f"{col}{r}"].font = F_LINK if col in "MNO" else F_TESTO
regole_mese = [
    ("J", "J5<85"),
    ("K", "K5<0.65"),
    ("L", "L5>0.05"),
    ("M", "M5<0.7"),
    ("N", "N5<0.7"),
    ("O", "O5<0.6"),
    ("I", "I5>1"),
    ("P", "P5<0.6"),
    ("Q", "Q5>15"),
    ("R", "R5>72"),
    ("S", "S5>1.1"),
    ("T", "T5>8"),
    ("U", "U5<7"),
]
for col, condizione in regole_mese:
    ms.conditional_formatting.add(f"{col}5:{col}16", FormulaRule(formula=[f"AND(ISNUMBER({col}5),{condizione})"], fill=ROSSO))
ms["A18"] = "Soglie (Parte 15.2-15.3): ricavo medio sotto 85 €; margine sotto 65%; abbandono sopra 5% (segnale se per due mesi); clienti e ricorrente sotto il 70% dell'obiettivo; cassa sotto il 60%; più di una fattura scaduta; risolte sotto 60%; più di 15 domande aperte; correzioni oltre 72 ore; minuti oltre 110%; attivazione oltre 8 giorni; soddisfazione sotto 7."
ms["A18"].font = F_NOTA

# ---------------------------------------------------------------------------
# Clienti (Parte 15.4 e 13.5)

cl = wb.create_sheet("Clienti")
titolo(cl, "Clienti", "Una riga per cliente attivo. Canone e minuti inclusi arrivano dal Listino. Rosso: minuti oltre il 110% o rischio alto (Parte 13.6).")
intestazione(
    cl,
    4,
    ["Locale", "Pacchetto", "Data di attivazione", "Canone (€)", "Minuti inclusi", "Minuti usati nel mese", "Uso dei minuti", "Mesi consecutivi oltre limite", "Ultimo contatto", "Giorni dall'ultimo contatto", "Rischio di abbandono", "Prossima azione", "Note"],
    [26, 11, 13, 11, 11, 13, 11, 13, 13, 13, 13, 34, 30],
)
for r in range(5, 45):
    for col in "ABCFHIKLM":
        cl[f"{col}{r}"].fill = DA_COMPILARE
        cl[f"{col}{r}"].font = F_TESTO
    cl[f"C{r}"].number_format = DATA
    cl[f"I{r}"].number_format = DATA
    cl[f"D{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(Listino!$C$5:$C$7,MATCH(B{r},Listino!$A$5:$A$7,0)),""))'
    cl[f"D{r}"].number_format = EURO
    cl[f"E{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(Listino!$E$5:$E$7,MATCH(B{r},Listino!$A$5:$A$7,0)),""))'
    cl[f"G{r}"] = f'=IF(OR(COUNT(F{r})=0,COUNT(E{r})=0),"",IF(E{r}=0,"",F{r}/E{r}))'
    cl[f"G{r}"].number_format = PERC
    cl[f"J{r}"] = f'=IF(AND(ISNUMBER(I{r}),I{r}<=TODAY()),TODAY()-I{r},"")'
    for col in "DEG":
        cl[f"{col}{r}"].font = F_LINK
    cl[f"J{r}"].font = F_TESTO
esempio_cliente = {"A": "Trattoria Esempio (riga di prova)", "B": "Pro", "C": dt.date(2026, 12, 11), "F": 180, "H": 0, "I": dt.date(2026, 12, 18), "K": "basso", "L": "Chiamata del giorno 30", "M": "Cliente fondatore"}
for col, v in esempio_cliente.items():
    cl[f"{col}5"] = v
dv_p2 = DataValidation(type="list", formula1='"Start,Pro,Premium"', allow_blank=True)
dv_rischio = DataValidation(type="list", formula1='"basso,medio,alto"', allow_blank=True)
cl.add_data_validation(dv_p2)
cl.add_data_validation(dv_rischio)
dv_p2.add("B5:B44")
dv_rischio.add("K5:K44")
cl.conditional_formatting.add("G5:G44", FormulaRule(formula=["AND(ISNUMBER(G5),G5>1.1)"], fill=ROSSO))
cl.conditional_formatting.add("H5:H44", FormulaRule(formula=["AND(ISNUMBER(H5),H5>=2)"], fill=ROSSO))
cl.conditional_formatting.add("K5:K44", CellIsRule(operator="equal", formula=['"alto"'], fill=ROSSO))
cl.conditional_formatting.add("J5:J44", FormulaRule(formula=["AND(ISNUMBER(J5),J5>30)"], fill=AMBRA))

# ---------------------------------------------------------------------------
# Concorrenti (Parte 7.4 e prompt 17.10)

co = wb.create_sheet("Concorrenti")
titolo(co, "Concorrenti", "Ricerca del 24/10/2026 (docs/ricerca/concorrenti.md). Aggiornare ogni trimestre e aggiungere quelli locali.")
intestazione(co, 4, ["Nome", "Cosa offre", "Prezzo", "Durata contratto", "Punto debole", "Fonte", "Verificato il"], [22, 40, 46, 16, 44, 46, 13])
concorrenti = [
    ("Fonio", "Assistente telefonico AI, anche ristoranti", "Solo 119 €/mese (99 € annuale) 1.000 minuti; Team 359 €/mese; extra 0,10-0,12 €/min", "Non indicata", "Da configurare da sé, nessuna persona locale", "https://www.fonio.ai/it/prezzi/"),
    ("CloudTalk AI Receptionist", "Receptionist AI su centralino cloud", "Da 99 €/mese per 200 minuti + 25 €/utente; extra 0,50 €/min", "Non verificata", "Costo per minuto alto, generico", "https://www.cloudtalk.io/it/blog/quanto-costa-servizio-risposta-ia/"),
    ("Retell AI", "Piattaforma per agenti vocali", "0,07-0,31 $/min a consumo", "Nessun minimo", "Strumento per sviluppatori", "https://www.retellai.com/pricing"),
    ("ElevenLabs Agents", "Piattaforma vocale (nostro fornitore)", "0,08 $/min; 15 minuti gratis", "A consumo", "Fornitore, non concorrente diretto", "https://elevenlabs.io/pricing/api"),
    ("Ainora", "Receptionist AI per ristoranti", "Non pubblicato", "Non indicata", "Dati elaborati anche negli USA di default", "https://ainora.lt/it/receptionist-ai-ristoranti"),
    ("Chiama.ai", "Centralino AI (Verona)", "Non pubblicato", "Non indicata", "Dichiarazione AI dal primo contatto: non verificato", "https://www.chiama.ai"),
    ("Cosmio", "Automazioni; receptionist AI open source", "Non pubblicato", "Non indicata", "Il receptionist si presenta con nome da persona", "https://cosmio.it"),
    ("TheFork Manager", "Prenotazioni + portale", "Non pubblicato in Italia; stima indiretta Francia ~139 €/mese + 2-4 €/coperto", "Non verificata", "Commissione per coperto", "https://restaurant.eatapp.co/en/blog/best-restaurant-reservation-software-france"),
    ("Zenchef", "Prenotazioni, sito, CRM", "Non pubblicato; indiretto ~150 €/mese medio", "Non verificata", "Non risponde al telefono", "https://www.businesscoot.com/en/company/zenchef"),
    ("CoverManager", "Prenotazioni e CRM", "Non pubblicato", "Non verificata", "Pensato anche per catene", "https://www.capterra.it/software/1012086/covermanager"),
    ("TeamSystem Cassa in Cloud", "Cassa e gestionale", "Non visibile; TS Scontrini + RT 30 € + IVA/mese annuale", "Annuale (TS Scontrini)", "Non risponde al telefono", "https://www.teamsystem.com/horeca/cassa-in-cloud/prezzi/"),
]
for i, riga in enumerate(concorrenti, start=5):
    for j, v in enumerate(riga, start=1):
        c = co.cell(row=i, column=j, value=v)
        c.font = F_TESTO
        c.alignment = Alignment(wrap_text=True, vertical="top")
    co.cell(row=i, column=7, value=dt.date(2026, 10, 24)).number_format = DATA
for r in range(16, 26):
    for col in range(1, 8):
        co.cell(row=r, column=col).fill = DA_COMPILARE

for foglio in wb.worksheets:
    foglio.sheet_view.showGridLines = foglio.title not in ("Leggimi",)

DESTINAZIONE.parent.mkdir(parents=True, exist_ok=True)
wb.save(DESTINAZIONE)
print(f"Scritto {DESTINAZIONE}")
